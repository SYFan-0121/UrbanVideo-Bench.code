import argparse
import os
import re
from pathlib import Path

import pandas as pd


QUESTION_CATEGORIES = [
    "Action Generation",
    "Association Reasoning",
    "Causal",
    "Cognitive Map",
    "Counterfactual",
    "Duration",
    "Goal Detection",
    "High-level Planning",
    "Landmark Position",
    "Object Recall",
    "Progress Evaluation",
    "Proximity",
    "Scene Recall",
    "Sequence Recall",
    "Start/End Position",
    "Trajectory Captioning",
]


def extract_option_letter(text):
    """Extract A-G from model output such as 'Option: [G]'."""
    if not isinstance(text, str):
        return None

    patterns = [
        r"Option\s*:\s*\[?\s*([A-Ga-g])\s*\]?",
        r"\[\s*([A-Ga-g])\s*\]",
        r"\b([A-Ga-g])\b",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(1).upper()
    return None


def find_latest_result_file(result_dir):
    candidates = sorted(
        Path(result_dir).glob("*_output.csv"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise FileNotFoundError(f"No *_output.csv files found in {result_dir}")
    return candidates[0]


def build_evaluation_tables(df):
    required = {"question_category", "answer", "Output"}
    missing = sorted(required - set(df.columns))
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    results = df.copy()
    results["GroundTruth"] = results["answer"].astype(str).str.strip().str.upper()
    results["Extracted_Option"] = results["Output"].apply(extract_option_letter)
    results["Is_Evaluated"] = results["Extracted_Option"].notna()
    results["Is_Correct"] = results["Extracted_Option"].eq(results["GroundTruth"])

    rows = []
    for category in QUESTION_CATEGORIES:
        group = results[results["question_category"] == category]
        total = int(len(group))
        evaluated = int(group["Is_Evaluated"].sum()) if total else 0
        correct = int(group["Is_Correct"].sum()) if total else 0
        rows.append(
            {
                "Category": category,
                "Total": total,
                "Evaluated": evaluated,
                "Correct": correct,
                "Accuracy": correct / evaluated if evaluated else 0.0,
                "Unanswered_or_Invalid": total - evaluated,
            }
        )

    category_df = pd.DataFrame(rows)
    total_questions = int(len(results))
    total_evaluated = int(results["Is_Evaluated"].sum())
    total_correct = int(results["Is_Correct"].sum())
    total_accuracy = total_correct / total_evaluated if total_evaluated else 0.0

    summary_df = pd.DataFrame(
        [
            {"Metric": "Total Questions", "Value": total_questions},
            {"Metric": "Evaluated Questions", "Value": total_evaluated},
            {"Metric": "Correct Answers", "Value": total_correct},
            {"Metric": "Overall Accuracy", "Value": total_accuracy},
            {"Metric": "Invalid or Missing Outputs", "Value": total_questions - total_evaluated},
        ]
    )

    detail_columns = [
        col
        for col in [
            "Question_id",
            "video_id",
            "question_category",
            "question",
            "answer",
            "Output",
            "Extracted_Option",
            "Is_Correct",
            "RetrievalStrategy",
            "RunID",
            "STRIVE",
        ]
        if col in results.columns
    ]
    detail_df = results[detail_columns].copy()
    return summary_df, category_df, detail_df


def autosize_columns(worksheet, max_width=80):
    for column_cells in worksheet.columns:
        letter = column_cells[0].column_letter
        width = 10
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            width = max(width, min(max_width, len(value) + 2))
        worksheet.column_dimensions[letter].width = width


def style_workbook(output_path):
    from openpyxl import load_workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(output_path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        autosize_columns(ws)

    summary_ws = wb["Summary"]
    for row in range(2, summary_ws.max_row + 1):
        if summary_ws.cell(row=row, column=1).value == "Overall Accuracy":
            summary_ws.cell(row=row, column=2).number_format = "0.00%"

    category_ws = wb["Category Accuracy"]
    accuracy_col = None
    for col in range(1, category_ws.max_column + 1):
        if category_ws.cell(row=1, column=col).value == "Accuracy":
            accuracy_col = col
            break

    if accuracy_col:
        for row in range(2, category_ws.max_row + 1):
            category_ws.cell(row=row, column=accuracy_col).number_format = "0.00%"

        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Accuracy by Question Category"
        chart.y_axis.title = "Accuracy"
        chart.x_axis.title = "Category"
        chart.height = 10
        chart.width = 24

        data = Reference(
            category_ws,
            min_col=accuracy_col,
            min_row=1,
            max_row=category_ws.max_row,
        )
        cats = Reference(
            category_ws,
            min_col=1,
            min_row=2,
            max_row=category_ws.max_row,
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.y_axis.scaling.min = 0
        chart.y_axis.scaling.max = 1
        category_ws.add_chart(chart, f"{get_column_letter(category_ws.max_column + 2)}2")

    wb.save(output_path)


def default_output_path(input_path):
    path = Path(input_path)
    stem = path.stem
    if stem.endswith("_output"):
        stem = stem[: -len("_output")]
    return path.with_name(f"{stem}_accuracy.xlsx")


def parse_args():
    parser = argparse.ArgumentParser(
        description="Convert benchmark result CSV into category and overall accuracy Excel report."
    )
    parser.add_argument(
        "file",
        nargs="?",
        help="Result CSV path. If omitted, the latest result/*_output.csv is used.",
    )
    parser.add_argument(
        "-o",
        "--output",
        help="Output Excel path. Defaults to <input>_accuracy.xlsx.",
    )
    parser.add_argument(
        "--result-dir",
        default="result",
        help="Directory used when auto-selecting the latest *_output.csv.",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    input_path = Path(args.file) if args.file else find_latest_result_file(args.result_dir)
    output_path = Path(args.output) if args.output else default_output_path(input_path)

    if not input_path.exists():
        raise FileNotFoundError(f"Result file not found: {input_path}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    df = pd.read_csv(input_path)
    summary_df, category_df, detail_df = build_evaluation_tables(df)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, index=False, sheet_name="Summary")
        category_df.to_excel(writer, index=False, sheet_name="Category Accuracy")
        detail_df.to_excel(writer, index=False, sheet_name="Question Results")

    style_workbook(output_path)
    print(f"Processed {input_path}")
    print(f"Saved Excel report to {output_path}")


if __name__ == "__main__":
    main()
